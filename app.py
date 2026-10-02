import streamlit as st
from langchain_community.document_loaders import (
    PyPDFLoader,
    Docx2txtLoader,
    TextLoader,
    WebBaseLoader,
    DirectoryLoader
)
from langchain.embeddings import HuggingFaceEmbeddings
from langchain_community.vectorstores import FAISS
from langchain.text_splitter import RecursiveCharacterTextSplitter
from langchain.chains.summarize import load_summarize_chain
from langchain.prompts import PromptTemplate
from dotenv import load_dotenv
import ollama
import tempfile
import os
from pathlib import Path
from typing import Optional, Tuple, List
from langchain_core.documents import Document
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
import re
from typing import List, Optional, Tuple, Dict
# Aşağıdaki importlar eksik
from datetime import datetime

from pptx.enum.shapes import MSO_SHAPE



# Çevre değişkenlerini yükle
load_dotenv()

# Hugging Face embedding modeli
embeddings = HuggingFaceEmbeddings(model_name="sentence-transformers/all-MiniLM-L6-v2")

def init_session_state():
    """Session state değişkenlerini başlat"""
    if "messages" not in st.session_state:
        st.session_state.messages = [{"role": "assistant", "content": "Nasıl yardımcı olabilirim?"}]
    if "vectorstore" not in st.session_state:
        st.session_state.vectorstore = None
    if "document_summaries" not in st.session_state:
        st.session_state.document_summaries = {}
    if "using_embedded_docs" not in st.session_state:
        st.session_state.using_embedded_docs = False

def load_document(file) -> Optional[List[Document]]:
    """Belge yükleme fonksiyonu"""
    try:
        file_extension = Path(file.name).suffix.lower()
        
        with tempfile.NamedTemporaryFile(delete=False) as temp_file:
            temp_file.write(file.getvalue())
            temp_file_path = temp_file.name

        if file_extension == '.pdf':
            loader = PyPDFLoader(temp_file_path)
        elif file_extension == '.docx':
            loader = Docx2txtLoader(temp_file_path)
        elif file_extension == '.txt':
            loader = TextLoader(temp_file_path)
        else:
            st.error(f"Desteklenmeyen dosya formatı: {file_extension}")
            return None
            
        documents = loader.load()
        return documents if documents else None

    except Exception as e:
        st.error(f"Dosya yükleme hatası: {str(e)}")
        return None
    finally:
        if os.path.exists(temp_file_path):
            os.unlink(temp_file_path)

def load_from_url(url: str) -> Optional[List[Document]]:
    """URL'den belge yükleme"""
    try:
        loader = WebBaseLoader(url)
        documents = loader.load()
        return documents if documents else None
    except Exception as e:
        st.error(f"Web sayfası yükleme hatası: {str(e)}")
        return None

def load_embedded_documents() -> Optional[List[Document]]:
    """Gömülü belgeleri yükle"""
    try:
        documents = []
        embedded_docs_path = os.path.join(os.path.dirname(__file__), 'belgeler')
        
        if not os.path.exists(embedded_docs_path):
            st.error("'belgeler' klasörü bulunamadı!")
            return None
            
        # PDF Loader
        pdf_loader = DirectoryLoader(
            embedded_docs_path,
            glob="**/*.pdf",
            loader_cls=PyPDFLoader
        )
        
        # DOCX Loader
        docx_loader = DirectoryLoader(
            embedded_docs_path,
            glob="**/*.docx",
            loader_cls=Docx2txtLoader
        )
        
        # TXT Loader
        txt_loader = DirectoryLoader(
            embedded_docs_path,
            glob="**/*.txt",
            loader_cls=TextLoader
        )
        
        # Tüm belgeleri yükle
        for loader in [pdf_loader, docx_loader, txt_loader]:
            try:
                docs = loader.load()
                if docs:
                    documents.extend(docs)
            except Exception as e:
                st.warning(f"Bazı belgeler yüklenirken hata oluştu: {str(e)}")
                continue
        
        return documents if documents else None

    except Exception as e:
        st.error(f"Gömülü belgeler yüklenirken hata oluştu: {str(e)}")
        return None

def process_documents(documents: List[Document]) -> Optional[FAISS]:
    """Belgeleri işle ve vektör veritabanı oluştur"""
    try:
        text_splitter = RecursiveCharacterTextSplitter(
            chunk_size=1000,
            chunk_overlap=200,
            length_function=len
        )
        texts = text_splitter.split_documents(documents)
        if not texts:
            return None
            
        return FAISS.from_documents(texts, embeddings)
    except Exception as e:
        st.error(f"Belge işleme hatası: {str(e)}")
        return None

def get_relevant_context(vectorstore: FAISS, query: str) -> Optional[str]:
    """Sorguyla ilgili bağlamı al"""
    try:
        docs = vectorstore.similarity_search(query)
        return "\n".join([doc.page_content for doc in docs])
    except Exception as e:
        st.error(f"Bağlam alma hatası: {str(e)}")
        return None




def summarize_document(document: Document) -> str:
    """Belge özeti oluştur"""
    try:
        prompt_template = """
        Aşağıdaki metni özetle:
        {text}
        
        Özet:
        """
        
        prompt = PromptTemplate(template=prompt_template, input_variables=["text"])
        chain = load_summarize_chain(
            llm=ollama,
            chain_type="stuff",
            prompt=prompt
        )
        
        summary = chain.run([document])
        return summary
    except Exception as e:
        return f"Özet oluşturma hatası: {str(e)}"
    


def format_text_shape(shape_or_paragraph, text: str, size: int = 18, color: RGBColor = None):
    """Metin biçimlendirmesi yapar"""
    # Check if it's a shape (has text_frame attribute) or a paragraph
    if hasattr(shape_or_paragraph, 'text_frame'):
        text_frame = shape_or_paragraph.text_frame
    else:
        text_frame = shape_or_paragraph
    
    if hasattr(text_frame, 'text'):
        text_frame.text = text
    
    # Get the paragraph (either the first paragraph of a text_frame or the paragraph itself)
    if hasattr(text_frame, 'paragraphs'):
        paragraph = text_frame.paragraphs[0]
    else:
        paragraph = text_frame
    
    # Apply formatting
    if hasattr(paragraph, 'font'):
        paragraph.font.size = Pt(size)
        if color:
            paragraph.font.color.rgb = color
        paragraph.font.name = "Calibri"
    
    if hasattr(paragraph, 'alignment'):
        paragraph.alignment = PP_ALIGN.LEFT

def generate_presentation_content(query: str, vectorstore: FAISS) -> dict:
    """Kullanıcı sorusuna ve kaynaklara göre sunum içeriği oluşturur"""
    try:
        # İlgili dokümanları al
        relevant_docs = vectorstore.similarity_search(query, k=5)
        context = "\n".join([doc.page_content for doc in relevant_docs])
        
        # Ollama ile içerik oluştur
        content_prompt = f"""
        Aşağıdaki konu ve kaynakları kullanarak kapsamlı bir sunum içeriği oluştur:
        
        KONU: {query}
        
        KAYNAKLAR:
        {context}
        
        Lütfen şu formatta yanıt ver:
        
        # GİRİŞ
        - [giriş madde 1]
        - [giriş madde 2]
        - [giriş madde 3]
        
        # ANA BÖLÜM
        - [ana bölüm madde 1]
        - [ana bölüm madde 2]
        - [ana bölüm madde 3]
        - [ana bölüm madde 4]
        
        # DETAYLAR
        - [detay madde 1]
        - [detay madde 2]
        - [detay madde 3]
        
        # SONUÇ
        - [sonuç madde 1]
        - [sonuç madde 2]
        - [sonuç madde 3]
        
        Her madde kaynaklardan alınan bilgilere dayanmalı ve en fazla bir cümle olmalı.
        """
        
        response = ollama.chat(
            model="llama3.1",
            messages=[
                {"role": "system", "content": "Sen bir sunum içeriği uzmanısın."},
                {"role": "user", "content": content_prompt}
            ]
        )
        
        return {
            'content': response['message']['content'],
            'documents': relevant_docs,
            'title': query
        }
    except Exception as e:
        raise Exception(f"Sunum içeriği oluşturma hatası: {str(e)}")

def create_presentation(title: str, content: str, relevant_docs: List[Document]) -> str:
    """PowerPoint sunumu oluşturur"""
    try:
        prs = Presentation()
        
        # Renk şeması
        colors = {
            'background': RGBColor(255, 255, 255),
            'primary': RGBColor(31, 73, 125),
            'secondary': RGBColor(149, 179, 215),
            'text': RGBColor(0, 0, 0)
        }

        # İçeriği bölümlere ayır
        sections = parse_presentation_content(content)

        # Kapak slaydı
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        title_shape = slide.shapes.title
        subtitle_shape = slide.placeholders[1]
        
        format_text_shape(title_shape, title, size=44, color=colors['primary'])
        format_text_shape(subtitle_shape, f"Hazırlanma Tarihi: {datetime.now().strftime('%d.%m.%Y')}", 
                         size=24, color=colors['text'])
        apply_slide_design(slide, colors, is_title_slide=True)

        # İçindekiler slaydı
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        title_shape = slide.shapes.title
        body_shape = slide.placeholders[1]
        
        format_text_shape(title_shape, "İçindekiler", size=36, color=colors['primary'])
        content_text = "\n".join([f"• {section['title']}" for section in sections])
        format_text_shape(body_shape, content_text, size=24, color=colors['text'])
        apply_slide_design(slide, colors)

        # Her bölüm için slayt oluştur
        for section in sections:
            slide = prs.slides.add_slide(prs.slide_layouts[1])
            title_shape = slide.shapes.title
            body_shape = slide.placeholders[1]
            
            format_text_shape(title_shape, section['title'], 
                            size=36, color=colors['primary'])
            content_text = "\n".join([f"• {point}" for point in section['points']])
            format_text_shape(body_shape, content_text, 
                            size=24, color=colors['text'])
            apply_slide_design(slide, colors)

        # Kaynakça slaydı
        create_sources_slide(prs, relevant_docs, colors)

        # Sunumu kaydet
        output_path = f"sunumlar/{slugify(title)}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pptx"
        os.makedirs(os.path.dirname(output_path), exist_ok=True)
        prs.save(output_path)
        return output_path

    except Exception as e:
        raise Exception(f"Sunum oluşturma hatası: {str(e)}")

def parse_presentation_content(content: str) -> List[dict]:
    """Sunum içeriğini bölümlere ayırır"""
    sections = []
    current_section = None
    
    for line in content.split('\n'):
        line = line.strip()
        if not line:
            continue
            
        if line.startswith('#'):
            if current_section:
                sections.append(current_section)
            current_section = {
                'title': line.lstrip('#').strip(),
                'points': []
            }
        elif line.startswith('-') and current_section:
            point = line.lstrip('- ').strip()
            if point:
                current_section['points'].append(point)
    
    if current_section:
        sections.append(current_section)
    
    return sections

def generate_response(prompt: str, context: str) -> str:
    """Yanıt üret"""
    try:
        st.session_state["full_message"] = ""
        
        # Sunum talebi kontrolü
        if any(keyword in prompt.lower() for keyword in ["sunum", "slayt", "powerpoint", "presentation"]):
            with st.spinner('Sunum hazırlanıyor...'):
                # Sunum içeriğini oluştur
                presentation_data = generate_presentation_content(prompt, st.session_state.vectorstore)
                
                # Sunumu oluştur
                presentation_path = create_presentation(
                    title=presentation_data['title'],
                    content=presentation_data['content'],
                    relevant_docs=presentation_data['documents']
                )
                
                # Sunum indirme butonu
                with open(presentation_path, "rb") as file:
                    st.download_button(
                        label="Sunumu İndir",
                        data=file,
                        file_name=f"{slugify(presentation_data['title'])}.pptx",
                        mime="application/vnd.openxmlformats-officedocument.presentationml.presentation"
                    )
                
                message = f"""Sunumunuz hazırlandı! 

İçerik kaynaklardan oluşturuldu ve şu bölümlerden oluşmaktadır:

1. Giriş
2. Ana Bölüm
3. Detaylar
4. Sonuç
5. Kaynakça

'Sunumu İndir' butonuna tıklayarak sunumu indirebilirsiniz.

İsterseniz sunum içeriğinde değişiklik yapmamı isteyebilirsiniz."""
                
                st.session_state["full_message"] = message
                yield message
                
        else:
            # Normal sohbet yanıtı
            response = ollama.chat(
                model="llama3.1",
                stream=True,
                messages=[
                    {"role": "system", "content": context},
                    {"role": "user", "content": prompt}
                ]
            )
            
            for partial_resp in response:
                token = partial_resp["message"]["content"]
                st.session_state["full_message"] += token
                yield token
            
    except Exception as e:
        error_message = f"Yanıt üretme hatası: {str(e)}"
        st.session_state["full_message"] = error_message
        yield error_message

def parse_summary(summary_text: str) -> List[dict]:
    """Özet metni bölümlere ayırır"""
    sections = []
    current_section = None
    
    for line in summary_text.split('\n'):
        line = line.strip()
        if not line:
            continue
            
        # Yeni bölüm başlangıcı
        if any(header in line.lower() for header in ['giriş', 'ana konu', 'detaylar', 'sonuç']):
            if current_section:
                sections.append(current_section)
            current_section = {'title': line, 'points': []}
        # Madde ekleme
        elif line.startswith('•') or line.startswith('-'):
            if current_section:
                point = line.lstrip('•- ').strip()
                current_section['points'].append(point)
    
    if current_section:
        sections.append(current_section)
    
    # Eğer bölüm yoksa varsayılan yapı oluştur
    if not sections:
        sections = [
            {'title': 'Giriş', 'points': ['Konu tanıtımı', 'Amaç ve hedefler']},
            {'title': 'Ana Konu', 'points': ['Temel kavramlar', 'Önemli noktalar']},
            {'title': 'Detaylar', 'points': ['Detaylı analiz', 'Örnekler']},
            {'title': 'Sonuç', 'points': ['Özet', 'Öneriler']}
        ]
    
    return sections

def get_document_reference(doc: Document) -> str:
    """Belge referansını formatla"""
    source = doc.metadata.get('source', 'Belirtilmemiş kaynak')
    return f"Kaynak: {source}"

def apply_slide_design(slide, colors: Dict[str, RGBColor], is_title_slide: bool = False):
    """Slayda tasarım uygular"""
    # Arka plan rengini ayarla
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = colors['background']
    
    # Kenar çizgisi ekle
    if not is_title_slide:
        # Slayt boyutlarını doğrudan presentation nesnesinden al
        prs = slide.part.package.presentation_part.presentation
        width = prs.slide_width
        height = prs.slide_height
        
        # Kenar boşluklarını ayarla
        margin = Inches(0.1)
        left = margin
        top = margin
        rect_width = width - (2 * margin)
        rect_height = height - (2 * margin)
        
        # Dikdörtgen şekli ekle
        shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, 
            left, 
            top, 
            rect_width, 
            rect_height
        )
        
        # Şekil özelliklerini ayarla
        shape.fill.background()
        shape.line.color.rgb = colors['secondary']
        shape.line.width = Pt(2)


def create_sources_slide(prs, relevant_docs: List[Document], colors: Dict[str, RGBColor]):
    """Kaynakça slaydı oluşturur"""
    sources_slide = prs.slides.add_slide(prs.slide_layouts[1])
    apply_slide_design(sources_slide, colors)
    
    format_text_shape(sources_slide.shapes.title, "Kaynakça", 
                     size=36, color=colors['primary'])
    
    content_shape = sources_slide.shapes.placeholders[1]
    for doc in relevant_docs:
        paragraph = content_shape.text_frame.add_paragraph()
        source_text = get_document_reference(doc)
        format_text_shape(paragraph, f"• {source_text}", 
                         size=20, color=colors['text'])
    
    return sources_slide

def slugify(text: str) -> str:
    """Metni dosya adı olarak kullanılabilir hale getirir"""
    text = text.lower()
    text = re.sub(r'[^\w\s-]', '', text)
    text = re.sub(r'[-\s]+', '-', text)
    return text.strip('-')




def analyze_content(content: str, relevant_docs: List[Document]) -> List[dict]:
    """İçeriği analiz edip slayt bölümlerine ayırır"""
    try:
        # Ollama modelini kullanarak içerik analizi yap
        analysis_prompt = f"""
        Aşağıdaki içeriği bir sunum için analiz et ve şu formatta döndür:
        - Her bölüm için bir başlık
        - Her başlık altında 3-5 arası madde
        - Maddeler kısa ve öz olmalı

        İçerik:
        {content}

        İlgili dökümanlar:
        {' '.join([doc.page_content for doc in relevant_docs])}
        """
        
        response = ollama.chat(model="llama3.1", messages=[
            {"role": "system", "content": "Sen bir sunum içeriği analisti ve düzenleyicisisin."},
            {"role": "user", "content": analysis_prompt}
        ])

        # Yanıtı işle ve bölümlere ayır
        sections = []
        current_section = None
        
        for line in response['message']['content'].split('\n'):
            line = line.strip()
            if not line:
                continue
                
            if line.startswith('#') or line.startswith('Bölüm'):
                if current_section:
                    sections.append(current_section)
                current_section = {'title': line.lstrip('#').strip(), 'points': []}
            elif line.startswith('-') or line.startswith('•'):
                if current_section:
                    current_section['points'].append(line.lstrip('-•').strip())
        
        if current_section:
            sections.append(current_section)
        
        # En az 3 bölüm olduğundan emin ol
        if len(sections) < 3:
            default_sections = [
                {'title': 'Giriş', 'points': ['Konu tanıtımı', 'Amaç ve hedefler', 'Genel bakış']},
                {'title': 'Ana Bölüm', 'points': ['Temel kavramlar', 'Önemli noktalar', 'Detaylı analiz']},
                {'title': 'Sonuç', 'points': ['Özet', 'Çıkarımlar', 'Öneriler']}
            ]
            sections.extend(default_sections[len(sections):])
        
        return sections
        
    except Exception as e:
        raise Exception(f"İçerik analizi hatası: {str(e)}")






def extract_main_points(content: str) -> List[str]:
    """Ana başlıkları çıkar"""
    # Başlıkları belirle (örnek: büyük harfle başlayan veya sayı ile başlayan satırlar)
    lines = content.split('\n')
    main_points = []
    
    for line in lines:
        line = line.strip()
        if line and (line[0].isupper() or line[0].isdigit()):
            main_points.append(line)
    
    # Eğer başlık bulunamazsa, içeriği mantıklı bölümlere ayır
    if not main_points:
        main_points = ["Giriş", "Ana Bölüm", "Sonuç"]
    
    return main_points

def extract_sub_points(title: str, content: str, relevant_docs: List[Document]) -> List[str]:
    """Her başlık için alt maddeleri belirle"""
    sub_points = []
    
    # İlgili belgelerden bu başlıkla alakalı bilgileri çıkar
    relevant_content = ""
    for doc in relevant_docs:
        if title.lower() in doc.page_content.lower():
            relevant_content += doc.page_content + "\n"
    
    # Alt maddeleri belirle
    sentences = re.split(r'[.!?]+', relevant_content)
    for sentence in sentences:
        sentence = sentence.strip()
        if sentence and len(sentence) > 20:  # Çok kısa cümleleri ele
            sub_points.append(sentence)
    
    # Eğer alt madde bulunamazsa varsayılan içerik ekle
    if not sub_points:
        sub_points = [
            f"{title} ile ilgili temel bilgiler",
            f"{title} konusunda önemli noktalar",
            f"{title} hakkında detaylar"
        ]
    
    return sub_points

def get_document_reference(doc: Document) -> str:
    """Belge referansını formatla"""
    try:
        # Belge kaynağını al (dosya adı veya URL)
        source = doc.metadata.get('source', 'Belirtilmemiş kaynak')
        
        # Kaynak türüne göre formatla
        if source.startswith('http'):
            return f"Web Kaynağı: {source}"
        else:
            return f"Belge: {source}"
            
    except Exception:
        return "Belirtilmemiş kaynak"
    

def get_relevant_documents(vectorstore: FAISS, query: str) -> List[Document]:
    """Sorguyla ilgili belgeleri al"""
    try:
        return vectorstore.similarity_search(query)
    except Exception:
        return []


def main():
    st.title("KURUMSAL İŞ ASİSTANI")
    init_session_state()
    
    # Yan menü
    with st.sidebar:
        st.header("Kaynak Seçimi")
        
        # Gömülü belgeler seçeneği
        use_embedded = st.checkbox(
            "Sadece Gömülü Belgeleri Kullan",
            value=st.session_state.using_embedded_docs
        )
        
        if use_embedded:
            st.session_state.using_embedded_docs = True
            process_embedded = st.button("Gömülü Belgeleri İşle")
            
            if process_embedded:
                with st.spinner('Gömülü belgeler yükleniyor...'):
                    embedded_docs = load_embedded_documents()
                    if embedded_docs:
                        vectorstore = process_documents(embedded_docs)
                        if vectorstore:
                            st.session_state.vectorstore = vectorstore
                            st.success("Gömülü belgeler başarıyla işlendi!")
                    else:
                        st.error("Gömülü belge bulunamadı!")
        
        else:
            st.session_state.using_embedded_docs = False
            # Birden fazla URL için giriş
            urls = st.text_area(
                "Web sayfası URL'lerini girin (her satıra bir URL):",
                "https://www.aselsan.com/tr\nhttps://www.aselsan.com/tr/hakkimizda\nhttps://www.aselsan.com/tr/iletisim\n"
            )
            
            # Dosya yükleme
            uploaded_files = st.file_uploader(
                "veya belge yükleyin (PDF, DOCX, TXT)",
                accept_multiple_files=True
            )
            
            process_button = st.button("Kaynakları İşle")
            
            if process_button:
                all_documents = []
                
                # URL'lerden yükleme
                if urls.strip():
                    url_list = urls.strip().split("\n")
                    with st.spinner('Web sayfaları yükleniyor...'):
                        for url in url_list:
                            web_docs = load_from_url(url.strip())
                            if web_docs:
                                all_documents.extend(web_docs)
                                st.success(f"{url.strip()} yüklendi!")
                
                # Dosyalardan yükleme
                if uploaded_files:
                    with st.spinner('Belgeler yükleniyor...'):
                        for file in uploaded_files:
                            docs = load_document(file)
                            if docs:
                                all_documents.extend(docs)
                                st.success(f"{file.name} yüklendi!")
                
                # Vektör veritabanı oluşturma
                if all_documents:
                    with st.spinner('Kaynaklar işleniyor...'):
                        vectorstore = process_documents(all_documents)
                        if vectorstore:
                            st.session_state.vectorstore = vectorstore
                            st.success("Tüm kaynaklar başarıyla işlendi!")
    
    # Ana sohbet arayüzü
    st.divider()
    
    for message in st.session_state.messages:
        with st.chat_message(message["role"]):
            st.markdown(message["content"])
    
    if prompt := st.chat_input("Sorunuzu yazın..."):
        st.session_state.messages.append({"role": "user", "content": prompt})
        
        with st.chat_message("user"):
            st.markdown(prompt)
        
        with st.chat_message("assistant"):
            if st.session_state.vectorstore is None:
                if st.session_state.using_embedded_docs:
                    st.markdown("Lütfen 'Gömülü Belgeleri İşle' butonuna basın.")
                else:
                    st.markdown("Lütfen önce kaynak ekleyin ve 'Kaynakları İşle' butonuna basın.")
            else:
                context = get_relevant_context(st.session_state.vectorstore, prompt)
                if context:
                    final_prompt = f"""
                    Soru: {prompt}
                    
                    Verilen bilgiler: {context}
                    
                    Yalnızca verilen bilgileri kullanarak soruyu yanıtla.
                    Eğer yanıt verilen bilgilerde bulunmuyorsa, 'Bu konuda bilgi veremiyorum.' de.
                    Yanıtı açık ve anlaşılır bir Türkçe olarak ver.
                    """
                    
                    st.write_stream(generate_response(prompt, final_prompt))
                    st.session_state.messages.append(
                        {"role": "assistant", "content": st.session_state["full_message"]}
                    )
                else:
                    st.markdown("Üzgünüm, sorunuzla ilgili bilgi bulamadım.")

if __name__ == "__main__":
    main()